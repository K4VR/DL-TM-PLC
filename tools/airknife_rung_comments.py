#!/usr/bin/env python3
"""Build a Studio 5000 V32 rung-comment import from a GE Logicmaster printout."""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from tools.compare_lm90_printouts import Rung, parse_printout

PLACEHOLDER = re.compile(r"^comment$", re.I)


def format_printout_comment(raw_lines: list[str]) -> str:
    """Turn Logicmaster (* ... *) print lines into readable comment text."""
    lines_out: list[str] = []
    pending_break = False
    for raw in raw_lines:
        s = raw.lstrip()
        if s[:1] in {"|", "*"}:
            s = s[1:]
        s = s.strip()
        s = re.sub(r"^\(\*+\s?", "", s)
        s = re.sub(r"\s*\*+\)$", "", s)
        s = s.strip(" *")
        if not s or set(s) <= {"*", "=", "-", "(", ")"}:
            pending_break = True
            continue
        if PLACEHOLDER.match(s):
            continue
        if pending_break and lines_out:
            lines_out.append("")
        lines_out.append(s)
        pending_break = False
    while lines_out and lines_out[0] == "":
        lines_out.pop(0)
    while lines_out and lines_out[-1] == "":
        lines_out.pop()
    return "\n".join(lines_out)


def logix_rung_comment(rung: Rung) -> str:
    ge = f"GE Rung {rung.number}"
    extra = format_printout_comment(rung.comments)
    if extra:
        return f"{ge}\n\n{extra}"
    return ge


def csv_comment(text: str) -> str:
    """Studio 5000 CSV uses $N for a newline inside a comment."""
    return text.replace("\r\n", "\n").replace("\n", "$N")


def routine_name(block: str, main_routine: str) -> str:
    return main_routine if block == "_MAIN" else block


def comment_rows(
    rungs: list[Rung],
    *,
    program: str,
    main_routine: str,
    location: str,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    index_by_block: dict[str, int] = {}
    for rung in rungs:
        seq = index_by_block.get(rung.block, 0)
        index_by_block[rung.block] = seq + 1
        loc = rung.number if location == "ge" else seq
        rows.append(
            {
                "program": program,
                "routine": routine_name(rung.block, main_routine),
                "block": rung.block,
                "ge_rung": rung.number,
                "logix_index": seq,
                "location": loc,
                "comment": logix_rung_comment(rung),
                "has_printout_comment": bool(format_printout_comment(rung.comments)),
            }
        )
    return rows


def write_csv(path: Path, rows: list[dict[str, object]], *, location_note: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(
            [
                "REMARK",
                "AIRKNIF 91026 rung comments for Studio 5000 Logix Designer V32. "
                "Tools > Import > Tags and Logic Comments. "
                "If Owning Element is blank, comments match by Location (rung number). "
                "Otherwise check 'Match all ladder diagram rung comments by rung number only'. "
                + location_note,
            ]
        )
        writer.writerow(
            [
                "REMARK",
                "Each comment starts with the GE printout rung number. "
                "Printout rung comments are appended when present. "
                "Change the program/routine names below if the Airknife project uses different names.",
            ]
        )
        for row in rows:
            writer.writerow(
                [
                    "RCOMMENT",
                    row["program"],
                    row["routine"],
                    csv_comment(str(row["comment"])),
                    "",
                    row["location"],
                ]
            )


def write_workbook(path: Path, sequential: list[dict[str, object]], ge_loc: list[dict[str, object]]) -> None:
    wb = Workbook()
    navy = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    title_font = Font(name="Calibri", bold=True, color="1F4E79", size=16)
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )

    ws = wb.active
    ws.title = "How to import"
    ws["A1"] = "Studio 5000 V32 rung comment import — AIRKNIF 91026"
    ws["A1"].font = title_font
    ws.merge_cells("A1:B1")
    steps = [
        ("CSV files", "Use the .csv next to this workbook. Sequential index is for a converted program whose Logix rungs are 0,1,2… in printout order. GE rung number is for a program whose Logix rung numbers still match the GE printout."),
        ("Open the Airknife project", "Studio 5000 Logix Designer V32, offline."),
        ("Import", "Tools → Import → Tags and Logic Comments. Choose the CSV."),
        ("Tag collisions", "Does not matter for this file (no TAG records)."),
        ("Logic comment collisions", "Import new comments and overwrite existing comments — unless you need to keep comments already typed in Logix."),
        ("Match by rung number", "If comments land on the wrong rungs, check Match all ladder diagram rung comments by rung number only. Owning Element is left blank so Location is used."),
        ("Program / routine names", "This file uses program Airknife, MainRoutine for GE _MAIN, and COAT_WG for the coating-weight block. Find/replace those names in the CSV if the V32 project differs."),
        ("Comment text", "Every imported comment starts with GE Rung N. When the 91026 printout had a rung comment, that text is added under the number."),
    ]
    ws["A3"] = "Step"
    ws["B3"] = "What to do"
    for col in (1, 2):
        ws.cell(3, col).fill = navy
        ws.cell(3, col).font = header_font
    for i, (step, detail) in enumerate(steps, 4):
        ws.cell(i, 1, step).alignment = wrap
        ws.cell(i, 2, detail).alignment = wrap
        ws.row_dimensions[i].height = 48
        for col in (1, 2):
            ws.cell(i, col).border = thin
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110

    def sheet(name: str, rows: list[dict[str, object]], loc_header: str) -> None:
        tab = wb.create_sheet(name)
        headers = ["Program", "Routine", "GE block", "GE Rung", loc_header, "Has printout comment", "Rung comment"]
        for col, val in enumerate(headers, 1):
            cell = tab.cell(1, col, val)
            cell.fill = navy
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = f"A1:G{1 + len(rows)}"
        for i, row in enumerate(rows, 2):
            values = [
                row["program"],
                row["routine"],
                row["block"],
                row["ge_rung"],
                row["location"],
                "Yes" if row["has_printout_comment"] else "No",
                row["comment"],
            ]
            for col, val in enumerate(values, 1):
                cell = tab.cell(i, col, val)
                cell.border = thin
                cell.alignment = wrap
            lines = str(row["comment"]).count("\n") + 1
            tab.row_dimensions[i].height = min(90, 15 * max(2, lines))
        widths = {1: 14, 2: 16, 3: 12, 4: 12, 5: 16, 6: 22, 7: 90}
        for col, width in widths.items():
            tab.column_dimensions[get_column_letter(col)].width = width

    sheet("Sequential Logix index", sequential, "Logix rung (0-based)")
    sheet("GE rung number", ge_loc, "Location = GE rung")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "printout",
        nargs="?",
        type=Path,
        default=Path("source/airknife/AIRKNIFE_2026-09-10.txt"),
        help="GE Logicmaster 91026 printout",
    )
    parser.add_argument("--program", default="Airknife", help="Logix program name (CSV Scope)")
    parser.add_argument("--main-routine", default="MainRoutine", help="Logix routine for GE _MAIN")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("import/airknife"),
        help="Directory for CSV and Excel files",
    )
    args = parser.parse_args(argv)

    parsed = parse_printout(args.printout, "91026")
    sequential = comment_rows(
        parsed.rungs, program=args.program, main_routine=args.main_routine, location="sequential"
    )
    ge_loc = comment_rows(
        parsed.rungs, program=args.program, main_routine=args.main_routine, location="ge"
    )
    out = args.output_dir
    seq_csv = out / "AIRKNIFE_91026_rung_comments_logix_index.csv"
    ge_csv = out / "AIRKNIFE_91026_rung_comments_ge_rung_number.csv"
    xlsx = out / "AIRKNIFE_91026_rung_comments.xlsx"
    write_csv(
        seq_csv,
        sequential,
        location_note="Location is the 0-based Logix rung index in printout order (GE _MAIN -> MainRoutine, then COAT_WG).",
    )
    write_csv(
        ge_csv,
        ge_loc,
        location_note="Location is the GE printout rung number (8, 11, 12, …), not a compacted 0-based index.",
    )
    write_workbook(xlsx, sequential, ge_loc)
    with_text = sum(1 for r in sequential if r["has_printout_comment"])
    print(f"Rungs: {len(sequential)}")
    print(f"With GE printout comment text: {with_text}")
    print(f"Wrote {seq_csv}")
    print(f"Wrote {ge_csv}")
    print(f"Wrote {xlsx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
